"""
Student files: solutions loaded from results/ and diagrams listed from results/diagrams/.
"""
import shutil
import subprocess
import tempfile
from pathlib import Path


_REPO_ROOT = Path(__file__).resolve().parents[3]
RESULTS_DIR = _REPO_ROOT / "results"
DIAGRAMS_DIR = RESULTS_DIR / "diagrams"

SUPPORTED_EXTENSIONS = (".md", ".txt", ".doc", ".docx")
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp")
_LIBREOFFICE_PATHS = (
    "soffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


class DocumentLoadError(Exception):
    pass


def load_solution_file(filename: str, results_dir: Path = RESULTS_DIR) -> tuple[Path, str]:
    """Loads a student file from results/ by name. Returns (path, text)."""
    name = filename.strip().strip('"').strip("'")
    path = results_dir / name

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise DocumentLoadError(
            f"'{name}' is not a supported format. Accepted: {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    if Path(name).name != name:
        raise DocumentLoadError(f"Use just the file name, without folders. The file must be in {results_dir.name}/.")
    if not path.is_file():
        available = list_result_files(results_dir)
        hint = f" Files there: {', '.join(available)}." if available else f" The {results_dir.name}/ folder is empty."
        raise DocumentLoadError(f"'{name}' was not found in {results_dir.name}/.{hint}")

    try:
        text = _read(path)
    except DocumentLoadError:
        raise
    except Exception as e:
        raise DocumentLoadError(f"Could not read '{name}': {type(e).__name__}: {e}") from e

    if not text.strip():
        raise DocumentLoadError(f"'{name}' is empty.")
    return path, text.strip()


def list_result_files(results_dir: Path = RESULTS_DIR) -> list[str]:
    if not results_dir.is_dir():
        return []
    return sorted(p.name for p in results_dir.iterdir() if p.suffix.lower() in SUPPORTED_EXTENSIONS)


def list_diagrams(diagrams_dir: Path = DIAGRAMS_DIR) -> list[Path]:
    if not diagrams_dir.is_dir():
        return []
    return sorted(p for p in diagrams_dir.iterdir() if p.is_file() and not p.name.startswith("."))


def _read(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return _read_text_file(path)
    if suffix == ".docx":
        return _read_docx(path)
    return _read_doc(path)


def _read_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError(f"Could not decode '{path.name}'. Save it as UTF-8 and try again.")


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        raise DocumentLoadError("Reading .docx needs python-docx: pip install python-docx") from None
    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)


def _read_doc(path: Path) -> str:
    # Legacy binary Word format — no pure-Python reader; convert with LibreOffice when available.
    soffice = next((p for p in _LIBREOFFICE_PATHS if shutil.which(p) or Path(p).is_file()), None)
    if not soffice:
        raise DocumentLoadError(
            f"Could not read '{path.name}': old .doc files need LibreOffice installed. "
            "Save it as .docx (File > Save As in Word) and try again."
        )
    with tempfile.TemporaryDirectory() as out_dir:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", out_dir, str(path)],
            check=True, capture_output=True, timeout=60,
        )
        return _read_text_file(Path(out_dir) / f"{path.stem}.txt")
